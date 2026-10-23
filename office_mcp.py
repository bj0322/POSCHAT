# SPDX-FileCopyrightText: 2026 Beomjoon Kim
# SPDX-License-Identifier: AGPL-3.0-only
"""내장 MCP 서버: Word(.docx)·PowerPoint(.pptx)·Excel(.xlsx) 문서를 만들고 고친다.

앱이 `PostechChat.exe --mcp-server word`(소스 실행은 `python office_mcp.py word`)로 띄우고,
표준 입출력으로 MCP(JSON-RPC 2.0, 줄 단위 JSON)를 주고받는다. 작업 폴더를 현재 폴더(cwd)로 받아
그 밖의 경로는 거부한다. 문서 읽기는 에이전트의 read 단계(documents.py)가 맡는다.

다른 내장 서버(paper_mcp·notion_mcp·firecrawl_mcp)도 이 모듈의 serve()로 돌고,
여기의 ToolFailure·경로 검사·HTTP 도우미를 함께 쓴다.
"""
import json
import os
import re
import sys
from pathlib import Path

PROTOCOL_VERSIONS = ("2025-06-18", "2025-03-26", "2024-11-05")
SERVER_VERSION = "1.0.0"


# 앱 프로세스 안의 스레드로 돌릴 때(안드로이드) 작업 폴더. None이면 현재 폴더(하위 프로세스로 실행할 때의 cwd).
ROOT = None
# initialize 때 받은 클라이언트 정보({"name", "version"}). memory 서버가 POSCHAT이 아닌 클라이언트를 거부하는 데 쓴다.
CLIENT_INFO = {}


class ToolFailure(Exception):
    pass


def _root():
    return Path(ROOT or Path.cwd()).resolve()


# ---- 경로 ----

def _path(path, suffixes):
    root = _root()
    p = (root / str(path)).resolve()
    if p != root and root not in p.parents:
        raise ToolFailure("작업 폴더 밖의 경로에는 쓸 수 없습니다.")
    if p.suffix.lower() not in suffixes:
        raise ToolFailure(f"파일 확장자는 {', '.join(suffixes)} 중 하나여야 합니다: {path}")
    return p


def _rel(p):
    return p.relative_to(_root()).as_posix()


def _new_file(path, suffixes, overwrite):
    p = _path(path, suffixes)
    if p.exists() and not overwrite:
        raise ToolFailure(f"이미 있는 파일입니다: {_rel(p)}. 덮어쓰려면 overwrite를 true로 주세요.")
    p.parent.mkdir(parents=True, exist_ok=True)
    return p


def _existing(path, suffixes):
    p = _path(path, suffixes)
    if not p.is_file():
        raise ToolFailure(f"파일이 없습니다: {_rel(p)}")
    return p


# ---- 네트워크(인터넷을 쓰는 내장 서버 공통) ----

_SESSION = None


def http(method, url, timeout=30, **kwargs):
    """requests 호출. 연결 실패는 ToolFailure로 바꾼다. 응답 상태는 호출한 쪽이 본다."""
    global _SESSION
    import requests

    if _SESSION is None:
        try:
            import truststore  # Windows 인증서 저장소를 쓴다(학교·회사 프록시 인증서). 안드로이드에는 없다

            truststore.inject_into_ssl()
        except ImportError:
            pass
        from version import APP_VERSION

        _SESSION = requests.Session()
        _SESSION.headers["User-Agent"] = f"PostechChat/{APP_VERSION} (+https://github.com/bj0322/POSCHAT)"
    try:
        return _SESSION.request(method, url, timeout=timeout, **kwargs)
    except requests.RequestException as e:
        raise ToolFailure(f"인터넷 요청에 실패했습니다({url.split('?')[0]}): {type(e).__name__}: {e}")


def _save(doc, p):
    try:
        doc.save(str(p))
    except PermissionError:
        raise ToolFailure(f"파일을 저장할 수 없습니다. 다른 프로그램(Office 등)에서 열려 있으면 닫고 다시 시도하세요: {_rel(p)}")


# ---- Word ----

INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*|_[^_\s][^_]*_)")


def _add_runs(paragraph, text):
    from docx.shared import Pt

    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif len(part) > 2 and part[0] == part[-1] and part[0] in "*_":
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _style(doc, name, fallback=None):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles[fallback] if fallback else None


def _add_markdown(doc, markdown):
    """간단한 마크다운을 Word 문단으로 옮긴다. 만든 블록 수를 돌려준다."""
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt

    lines = str(markdown).replace("\r\n", "\n").split("\n")
    count, i = 0, 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code))
            run.font.name = "Consolas"
            run.font.size = Pt(10)
            count += 1
            continue
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                ncols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=ncols)
                table.style = _style(doc, "Table Grid")
                for r, cells in enumerate(rows):
                    for c, value in enumerate(cells):
                        cell_p = table.cell(r, c).paragraphs[0]
                        _add_runs(cell_p, value)
                        if r == 0:
                            for run in cell_p.runs:
                                run.bold = True
                count += 1
            continue
        if stripped == "---":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif m := re.match(r"(#{1,6})\s+(.*)", stripped):
            doc.add_heading(m.group(2).strip(), level=len(m.group(1)))
        elif m := re.match(r"( *)[-*+]\s+(.*)", line):
            level = min(len(m.group(1)) // 2, 2)
            p = doc.add_paragraph(style=_style(doc, "List Bullet" + (f" {level + 1}" if level else ""), "List Bullet"))
            _add_runs(p, m.group(2))
        elif m := re.match(r"( *)\d+[.)]\s+(.*)", line):
            level = min(len(m.group(1)) // 2, 2)
            p = doc.add_paragraph(style=_style(doc, "List Number" + (f" {level + 1}" if level else ""), "List Number"))
            _add_runs(p, m.group(2))
        elif stripped.startswith(">"):
            p = doc.add_paragraph(style=_style(doc, "Quote", "Normal"))
            _add_runs(p, stripped.lstrip("> ").strip())
        else:
            _add_runs(doc.add_paragraph(), stripped)
        count += 1
        i += 1
    return count


def _korean_font(doc):
    from docx.oxml.ns import qn

    rpr = doc.styles["Normal"].element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), "맑은 고딕")


def word_create_document(path, markdown, title=None, overwrite=False):
    from docx import Document

    p = _new_file(path, (".docx",), overwrite)
    doc = Document()
    _korean_font(doc)
    if title:
        doc.add_heading(str(title), level=0)
        doc.core_properties.title = str(title)
    blocks = _add_markdown(doc, markdown)
    _save(doc, p)
    return f"Word 문서를 만들었습니다: {_rel(p)} (블록 {blocks}개)"


def word_append_content(path, markdown):
    from docx import Document

    p = _existing(path, (".docx",))
    doc = Document(str(p))
    blocks = _add_markdown(doc, markdown)
    _save(doc, p)
    return f"{_rel(p)} 끝에 블록 {blocks}개를 추가했습니다."


def _docx_paragraphs(doc):
    def walk(container):
        for paragraph in container.paragraphs:
            yield paragraph
        for table in getattr(container, "tables", []):
            for row in table.rows:
                for cell in row.cells:
                    yield from walk(cell)

    yield from walk(doc)
    for section in doc.sections:
        for part in (section.header, section.footer):
            yield from walk(part)


def _replace_in_paragraph(paragraph, old, new):
    """서식을 지키려고 run 안에서 바꾸고, old가 여러 run에 걸쳐 있으면 문단 글자를 첫 run에 모아 바꾼다."""
    runs = list(paragraph.runs)
    full = "".join(run.text for run in runs)
    total = full.count(old)
    if not total:
        return 0
    if sum(run.text.count(old) for run in runs) == total:
        for run in runs:
            run.text = run.text.replace(old, new)
    else:
        runs[0].text = full.replace(old, new)
        for run in runs[1:]:
            run.text = ""
    return total


def word_replace_text(path, old, new):
    from docx import Document

    if not old:
        raise ToolFailure("old가 비어 있습니다.")
    p = _existing(path, (".docx",))
    doc = Document(str(p))
    hits = sum(_replace_in_paragraph(par, str(old), str(new)) for par in _docx_paragraphs(doc))
    if not hits:
        raise ToolFailure("old 문자열을 문서에서 찾지 못했습니다. read로 문서 내용을 다시 확인하세요.")
    _save(doc, p)
    return f"{_rel(p)}에서 {hits}곳을 바꿨습니다."


# ---- PowerPoint ----

def _widen(slide, ratio):
    # 기본 서식 파일의 레이아웃은 4:3 기준이라 16:9로 넓힌 슬라이드에 맞게 가로 위치·폭을 늘린다.
    for shape in slide.placeholders:
        if shape.left is not None and shape.width is not None:
            shape.left, shape.width = int(shape.left * ratio), int(shape.width * ratio)


def _fill_bullets(frame, bullets):
    frame.clear()
    first = True
    for item in bullets:
        text = str(item)
        level = min((len(text) - len(text.lstrip(" "))) // 2, 4)
        text = re.sub(r"^\s*(?:[-*•]\s+)?", "", text)
        para = frame.paragraphs[0] if first else frame.add_paragraph()
        para.text, para.level = text, level
        first = False


def _add_slide(prs, spec, ratio):
    from pptx.util import Emu, Pt

    spec = spec if isinstance(spec, dict) else {"title": str(spec)}
    layout = str(spec.get("layout") or ("table" if spec.get("table") else "content")).lower()
    index = {"title": 0, "content": 1, "section": 2, "table": 5, "title_only": 5, "blank": 6}.get(layout, 1)
    slide = prs.slides.add_slide(prs.slide_layouts[index])
    _widen(slide, ratio)
    if slide.shapes.title is not None and spec.get("title"):
        slide.shapes.title.text = str(spec["title"])
    body = [ph for ph in slide.placeholders if ph.placeholder_format.idx == 1]
    if layout == "title" and body and spec.get("subtitle"):
        body[0].text = str(spec["subtitle"])
    elif body and spec.get("bullets"):
        _fill_bullets(body[0].text_frame, spec["bullets"])
    elif body and spec.get("text"):
        body[0].text = str(spec["text"])
    rows = spec.get("table")
    if rows:
        rows = [[str(c) for c in r] for r in rows]
        ncols = max(len(r) for r in rows)
        margin = Emu(int(prs.slide_width * 0.06))
        top = Emu(int(prs.slide_height * 0.25))
        shape = slide.shapes.add_table(len(rows), ncols, margin, top, prs.slide_width - margin * 2,
                                       Emu(int(prs.slide_height * 0.08) * len(rows)))
        for r, cells in enumerate(rows):
            for c, value in enumerate(cells):
                cell = shape.table.cell(r, c)
                cell.text = value
                for para in cell.text_frame.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(14)
    if spec.get("notes"):
        slide.notes_slide.notes_text_frame.text = str(spec["notes"])
    return slide


def _slides_arg(slides):
    if isinstance(slides, str):
        try:
            slides = json.loads(slides)
        except ValueError:
            raise ToolFailure("slides는 슬라이드 객체의 배열이어야 합니다.")
    if not isinstance(slides, list) or not slides:
        raise ToolFailure("slides는 슬라이드 객체가 하나 이상 있는 배열이어야 합니다.")
    return slides


def powerpoint_create_presentation(path, slides, widescreen=True, overwrite=False):
    from pptx import Presentation
    from pptx.util import Inches

    slides = _slides_arg(slides)
    p = _new_file(path, (".pptx",), overwrite)
    prs = Presentation()
    ratio = 1.0
    if widescreen:
        ratio = Inches(13.333) / prs.slide_width
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for spec in slides:
        _add_slide(prs, spec, ratio)
    _save(prs, p)
    return f"프레젠테이션을 만들었습니다: {_rel(p)} (슬라이드 {len(slides)}장)"


def powerpoint_add_slides(path, slides):
    from pptx import Presentation

    slides = _slides_arg(slides)
    p = _existing(path, (".pptx",))
    prs = Presentation(str(p))
    # 기존 파일의 레이아웃이 실제 슬라이드 크기와 다르면(4:3 서식을 16:9로 넓힌 경우) 같은 비율로 맞춘다.
    title = next((ph for ph in prs.slide_layouts[1].placeholders if ph.placeholder_format.idx == 0), None)
    ratio = 1.0
    if title is not None and title.left is not None and title.left + title.width < prs.slide_width * 0.8:
        ratio = prs.slide_width / max(title.left * 2 + title.width, 1)
    for spec in slides:
        _add_slide(prs, spec, ratio)
    _save(prs, p)
    return f"{_rel(p)}에 슬라이드 {len(slides)}장을 추가했습니다. (전체 {len(prs.slides)}장)"


def _pptx_frames(prs):
    def from_shapes(shapes):
        for shape in shapes:
            if shape.shape_type == 6:  # 그룹
                yield from from_shapes(shape.shapes)
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                yield shape.text_frame
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        yield cell.text_frame

    for slide in prs.slides:
        yield from from_shapes(slide.shapes)
        if slide.has_notes_slide:
            yield slide.notes_slide.notes_text_frame


def powerpoint_replace_text(path, old, new):
    from pptx import Presentation

    if not old:
        raise ToolFailure("old가 비어 있습니다.")
    p = _existing(path, (".pptx",))
    prs = Presentation(str(p))
    hits = sum(_replace_in_paragraph(para, str(old), str(new))
               for frame in _pptx_frames(prs) for para in frame.paragraphs)
    if not hits:
        raise ToolFailure("old 문자열을 프레젠테이션에서 찾지 못했습니다. read로 내용을 다시 확인하세요.")
    _save(prs, p)
    return f"{_rel(p)}에서 {hits}곳을 바꿨습니다."


# ---- Excel ----

def _rows_arg(rows):
    if isinstance(rows, str):
        try:
            rows = json.loads(rows)
        except ValueError:
            raise ToolFailure("rows는 행(값의 배열)의 배열이어야 합니다.")
    if not isinstance(rows, list) or not all(isinstance(r, list) for r in rows):
        raise ToolFailure("rows는 행(값의 배열)의 배열이어야 합니다. 예: [[\"이름\", \"점수\"], [\"홍길동\", 90]]")
    return rows


def _write_rows(ws, rows, start="A1", header=False):
    from openpyxl.styles import Font
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter

    try:
        col_letter, row0 = coordinate_from_string(str(start).upper())
    except ValueError:
        raise ToolFailure(f"시작 셀 주소가 올바르지 않습니다: {start}")
    col0 = column_index_from_string(col_letter)
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            cell = ws.cell(row=row0 + r, column=col0 + c)
            cell.value = value
            if header and r == 0:
                cell.font = Font(bold=True)
    end = f"{get_column_letter(col0 + max((len(r) for r in rows), default=1) - 1)}{row0 + max(len(rows), 1) - 1}"
    return f"{str(start).upper()}:{end}"


def _autofit(ws):
    from openpyxl.utils import get_column_letter

    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                text = str(cell.value)
                # 한글은 영문보다 폭이 넓다.
                width = sum(2 if ord(ch) > 0x2E80 else 1 for ch in text)
                widths[cell.column] = max(widths.get(cell.column, 0), width)
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 8), 60)


def excel_create_workbook(path, sheets, overwrite=False):
    from openpyxl import Workbook

    if isinstance(sheets, str):
        try:
            sheets = json.loads(sheets)
        except ValueError:
            raise ToolFailure("sheets는 시트 객체의 배열이어야 합니다.")
    if not isinstance(sheets, list) or not sheets:
        raise ToolFailure("sheets는 {\"name\", \"rows\"} 객체가 하나 이상 있는 배열이어야 합니다.")
    p = _new_file(path, (".xlsx",), overwrite)
    wb = Workbook()
    wb.remove(wb.active)
    summary = []
    for i, spec in enumerate(sheets, 1):
        spec = spec if isinstance(spec, dict) else {"rows": spec}
        ws = wb.create_sheet(str(spec.get("name") or f"Sheet{i}")[:31])
        rows = _rows_arg(spec.get("rows") or [])
        if rows:
            _write_rows(ws, rows, "A1", header=spec.get("header", True))
            _autofit(ws)
            ws.freeze_panes = "A2" if spec.get("header", True) and len(rows) > 1 else None
        summary.append(f"{ws.title}({len(rows)}행)")
    _save(wb, p)
    return f"Excel 파일을 만들었습니다: {_rel(p)} — " + ", ".join(summary)


def excel_write_cells(path, rows, sheet=None, start_cell="A1"):
    from openpyxl import Workbook, load_workbook

    rows = _rows_arg(rows)
    p = _path(path, (".xlsx", ".xlsm"))
    created = not p.exists()
    if created:
        p.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        if sheet:
            wb.active.title = str(sheet)[:31]
    else:
        wb = load_workbook(str(p), keep_vba=p.suffix.lower() == ".xlsm")
    if sheet and sheet not in wb.sheetnames:
        ws = wb.create_sheet(str(sheet)[:31])
    else:
        ws = wb[sheet] if sheet else wb.active
    area = _write_rows(ws, rows, start_cell)
    _save(wb, p)
    return f"{_rel(p)}의 '{ws.title}' 시트 {area}에 값을 썼습니다." + (" (새 파일)" if created else "")


# ---- 도구 정의 ----

def _schema(props, required):
    return {"type": "object", "properties": props, "required": required}


STR = {"type": "string"}
BOOL = {"type": "boolean"}
INT = {"type": "integer"}
# 도구 정의 (함수, 설명, 입력 스키마[, 주석]). 주석을 생략하면 파일을 고치는 도구로 본다.
READ_ONLY = {"readOnlyHint": True, "destructiveHint": False, "openWorldHint": True}
WRITES = {"readOnlyHint": False, "destructiveHint": False, "openWorldHint": False}
SLIDE = {
    "type": "object",
    "description": "layout: title|content|section|table|blank (기본 content). title, subtitle(title 레이아웃), "
                   "bullets(문자열 배열, 앞 공백 2칸마다 한 단계 들여쓰기), text, table(행 배열), notes(발표자 노트)",
}
MARKDOWN_HELP = ("마크다운: # 제목(1~6단계), - 글머리 기호(공백 2칸마다 들여쓰기), 1. 번호 목록, > 인용, "
                 "| 표 |, ``` 코드 ```, **굵게**, *기울임*, `코드`, ---만 있는 줄은 페이지 나누기")

SERVERS = {
    "word": {
        "create_document": (word_create_document, "새 Word 문서(.docx)를 마크다운 내용으로 만듭니다. " + MARKDOWN_HELP,
                            _schema({"path": STR, "markdown": STR, "title": STR, "overwrite": BOOL}, ["path", "markdown"])),
        "append_content": (word_append_content, "기존 Word 문서 끝에 마크다운 내용을 덧붙입니다.",
                           _schema({"path": STR, "markdown": STR}, ["path", "markdown"])),
        "replace_text": (word_replace_text, "Word 문서(본문·표·머리글·바닥글)에서 old를 모두 new로 바꿉니다. 서식은 최대한 유지합니다.",
                         _schema({"path": STR, "old": STR, "new": STR}, ["path", "old", "new"])),
    },
    "powerpoint": {
        "create_presentation": (powerpoint_create_presentation,
                                "새 PowerPoint(.pptx)를 만듭니다. slides는 슬라이드 객체 배열입니다. 첫 장은 보통 layout \"title\"로 합니다.",
                                _schema({"path": STR, "slides": {"type": "array", "items": SLIDE},
                                         "widescreen": BOOL, "overwrite": BOOL}, ["path", "slides"])),
        "add_slides": (powerpoint_add_slides, "기존 PowerPoint 끝에 슬라이드를 추가합니다.",
                       _schema({"path": STR, "slides": {"type": "array", "items": SLIDE}}, ["path", "slides"])),
        "replace_text": (powerpoint_replace_text, "모든 슬라이드·표·발표자 노트에서 old를 new로 바꿉니다.",
                         _schema({"path": STR, "old": STR, "new": STR}, ["path", "old", "new"])),
    },
    "excel": {
        "create_workbook": (excel_create_workbook,
                            "새 Excel(.xlsx)을 만듭니다. sheets는 {name, rows, header} 배열이고 rows는 행 배열입니다. "
                            "\"=SUM(B2:B5)\"처럼 =로 시작하는 문자열은 수식이 됩니다. header(기본 true)는 첫 행을 굵게 합니다.",
                            _schema({"path": STR, "sheets": {"type": "array", "items": {"type": "object"}},
                                     "overwrite": BOOL}, ["path", "sheets"])),
        "write_cells": (excel_write_cells,
                        "Excel 파일의 시트에 start_cell(기본 A1)부터 rows를 씁니다. 파일이나 시트가 없으면 만듭니다. "
                        "read로 본 셀 주소(예: B3)를 그대로 쓰면 그 셀을 고칩니다.",
                        _schema({"path": STR, "rows": {"type": "array", "items": {"type": "array"}},
                                 "sheet": STR, "start_cell": STR}, ["path", "rows"])),
    },
}


def all_servers():
    """내장 서버 이름 -> 도구 정의. 다른 서버 모듈은 이 모듈을 가져다 쓰므로 여기서 늦게 가져온다.

    github·google 등은 커넥터(connectors.py)를 켰을 때만 mcp_client가 띄운다.
    """
    import firecrawl_mcp
    import github_mcp
    import google_mcp
    import notion_mcp
    import paper_mcp
    import memory_mcp
    import research_mcp
    import runner_mcp

    return {**SERVERS, "paper": paper_mcp.TOOLS, "notion": notion_mcp.TOOLS, "firecrawl": firecrawl_mcp.TOOLS,
            "github": github_mcp.TOOLS, **google_mcp.SERVERS,
            "research": research_mcp.TOOLS, "zotero": research_mcp.ZOTERO_TOOLS, "memory": memory_mcp.TOOLS,
            "python": runner_mcp.PYTHON_TOOLS, "ssh": runner_mcp.SSH_TOOLS}


# ---- MCP 표준 입출력 서버 ----

def _std_streams():
    """창 모드 exe는 sys.stdin/stdout이 None일 수 있어서 상속받은 표준 핸들을 직접 연다."""
    if sys.stdin is not None and sys.stdout is not None:
        return sys.stdin.buffer, sys.stdout.buffer
    import ctypes
    import msvcrt

    get = ctypes.windll.kernel32.GetStdHandle
    get.restype = ctypes.c_void_p
    fin = os.fdopen(msvcrt.open_osfhandle(get(-10), os.O_RDONLY), "rb")
    fout = os.fdopen(msvcrt.open_osfhandle(get(-11), 0), "wb")
    return fin, fout


def _handle(name, tools, msg):
    method, params = msg.get("method"), msg.get("params") or {}
    if method == "initialize":
        global CLIENT_INFO
        CLIENT_INFO = params.get("clientInfo") if isinstance(params.get("clientInfo"), dict) else {}
        version = params.get("protocolVersion")
        return {
            "protocolVersion": version if version in PROTOCOL_VERSIONS else PROTOCOL_VERSIONS[0],
            "capabilities": {"tools": {"listChanged": False}},
            "serverInfo": {"name": f"postech-chat-{name}", "version": SERVER_VERSION},
        }
    if method == "ping":
        return {}
    if method == "tools/list":
        return {"tools": [{
            "name": tool, "description": entry[1], "inputSchema": entry[2],
            "annotations": entry[3] if len(entry) > 3 else WRITES,
        } for tool, entry in tools.items()]}
    if method == "tools/call":
        entry = tools.get(params.get("name"))
        if entry is None:
            raise LookupError(f"알 수 없는 도구: {params.get('name')}")
        try:
            text, is_error = entry[0](**(params.get("arguments") or {})), False
        except ToolFailure as e:
            text, is_error = f"오류: {e}", True
        except TypeError as e:
            text, is_error = f"오류: 잘못된 인자 ({e})", True
        except Exception as e:
            text, is_error = f"오류: {type(e).__name__}: {e}", True
        return {"content": [{"type": "text", "text": text}], "isError": is_error}
    raise NotImplementedError(method)


def serve(name, fin=None, fout=None):
    """MCP 서버를 돌린다. fin/fout을 주지 않으면 표준 입출력을 쓴다."""
    servers = all_servers()
    tools = servers.get(name)
    if tools is None:
        sys.exit(f"알 수 없는 MCP 서버: {name} (사용 가능: {', '.join(servers)})")
    if fin is None:
        fin, fout = _std_streams()
    for line in fin:
        line = line.strip()
        if not line:
            continue
        try:
            msg = json.loads(line)
        except ValueError:
            reply = {"jsonrpc": "2.0", "id": None, "error": {"code": -32700, "message": "Parse error"}}
        else:
            if "id" not in msg:
                continue  # notifications/initialized 등 알림은 답하지 않는다
            try:
                reply = {"jsonrpc": "2.0", "id": msg["id"], "result": _handle(name, tools, msg)}
            except NotImplementedError as e:
                reply = {"jsonrpc": "2.0", "id": msg["id"], "error": {"code": -32601, "message": f"Method not found: {e}"}}
            except LookupError as e:
                reply = {"jsonrpc": "2.0", "id": msg["id"], "error": {"code": -32602, "message": str(e)}}
        fout.write((json.dumps(reply, ensure_ascii=False) + "\n").encode("utf-8"))
        fout.flush()


if __name__ == "__main__":
    # 다른 서버 모듈이 `import office_mcp`로 가져가는 모듈과 같은 것을 써야 ToolFailure·ROOT가 하나가 된다.
    import office_mcp

    office_mcp.serve(sys.argv[1] if len(sys.argv) > 1 else "")
