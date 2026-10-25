# SPDX-FileCopyrightText: 2026 Beomjoon Kim
# SPDX-License-Identifier: AGPL-3.0-only
"""PDF·Word·PowerPoint·Excel·한글 문서에서 에이전트가 읽을 텍스트를 뽑는다.

read와 search는 이 모듈로 문서를 텍스트로 바꿔 다룬다. 문서를 만들거나 고치는 일은
내장 오피스 MCP 서버(office_mcp.py)가 맡는다.
"""
import os
import re
import struct
import zipfile
import zlib
from pathlib import Path
from xml.etree import ElementTree as ET

DOCUMENT_EXTS = {".pdf", ".docx", ".pptx", ".xlsx", ".xlsm", ".hwp", ".hwpx"}
# 예전 Office 바이너리 형식은 읽지 못하므로 안내만 한다.
LEGACY_EXTS = {".doc": ".docx", ".ppt": ".pptx", ".xls": ".xlsx"}
MAX_DOCUMENT_BYTES = 50 * 1024 * 1024
MAX_SHEET_ROWS = 2000
MAX_SHEET_COLS = 50

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


class DocumentError(Exception):
    pass


def is_document(path):
    suffix = os.path.splitext(str(path))[1].lower()   # 원격 경로(remote_fs.RPath, ssh://…)도 이름으로만 본다
    return suffix in DOCUMENT_EXTS or suffix in LEGACY_EXTS


def extract_text(path):
    """문서 파일의 텍스트. 지원하지 않거나 읽을 수 없으면 DocumentError."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in LEGACY_EXTS:
        raise DocumentError(f"예전 형식({suffix})은 읽을 수 없습니다. Office에서 {LEGACY_EXTS[suffix]}로 저장한 뒤 다시 시도하세요.")
    if path.stat().st_size > MAX_DOCUMENT_BYTES:
        raise DocumentError(f"문서가 너무 큽니다({MAX_DOCUMENT_BYTES // 1024 // 1024}MB 초과).")
    reader = {
        ".pdf": _pdf, ".docx": _docx, ".pptx": _pptx, ".xlsx": _xlsx, ".xlsm": _xlsx,
        ".hwp": _hwp, ".hwpx": _hwpx,
    }.get(suffix)
    if reader is None:
        raise DocumentError(f"지원하지 않는 문서 형식입니다: {suffix}")
    try:
        text = reader(path)
    except DocumentError:
        raise
    except (zipfile.BadZipFile, KeyError, ET.ParseError) as e:
        raise DocumentError(f"문서 구조를 읽지 못했습니다(손상되었거나 다른 형식): {e}")
    except Exception as e:
        raise DocumentError(f"문서를 읽지 못했습니다: {type(e).__name__}: {e}")
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return text or "(문서에서 텍스트를 찾지 못했습니다. 스캔한 이미지 문서일 수 있습니다.)"


# ---- PDF ----

def _pdf(path):
    from pypdf import PdfReader
    from pypdf.errors import FileNotDecryptedError

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise DocumentError("암호가 걸린 PDF입니다.")
    pages = []
    try:
        for i, page in enumerate(reader.pages, 1):
            # PDF는 글자 위치로 공백을 추정하므로 겹친 공백을 하나로 줄인다.
            text = re.sub(r"[ \t]{2,}", " ", page.extract_text() or "").strip()
            pages.append(f"--- 페이지 {i} ---\n{text}")
    except FileNotDecryptedError:
        raise DocumentError("암호가 걸린 PDF입니다.")
    return "\n\n".join(pages)


# ---- Word ----

def _docx_paragraph(p):
    parts = []
    for node in p.iter():
        if node.tag == W + "t":
            parts.append(node.text or "")
        elif node.tag == W + "tab":
            parts.append("\t")
        elif node.tag in (W + "br", W + "cr"):
            parts.append("\n")
    text = "".join(parts)
    style = p.find(f"{W}pPr/{W}pStyle")
    style = style.get(W + "val", "") if style is not None else ""
    m = re.match(r"(?:Heading|제목)\s*(\d)", style)
    if m and text.strip():
        return "#" * int(m.group(1)) + " " + text
    if style == "Title" and text.strip():
        return "# " + text
    if (style.startswith("List") or p.find(f"{W}pPr/{W}numPr") is not None) and text.strip():
        return "- " + text
    return text


def _docx_table(tbl):
    rows = []
    for tr in tbl.findall(W + "tr"):
        cells = []
        for tc in tr.findall(W + "tc"):
            cells.append(" ".join(_docx_paragraph(p).strip() for p in tc.iter(W + "p")).strip())
        rows.append("| " + " | ".join(cells) + " |")
    return "\n".join(rows)


def _docx(path):
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    body = root.find(W + "body")
    blocks = []
    for child in body if body is not None else []:
        if child.tag == W + "p":
            blocks.append(_docx_paragraph(child))
        elif child.tag == W + "tbl":
            blocks.append(_docx_table(child))
    return "\n".join(blocks)


# ---- PowerPoint ----

def _num(name):
    m = re.search(r"(\d+)\.xml$", name)
    return int(m.group(1)) if m else 0


def _a_text(p):
    return "".join(t.text or "" for t in p.iter(A + "t"))


def _pptx_notes(z, slide_name):
    """슬라이드의 rels에서 발표자 노트 파일을 찾는다(노트 번호는 슬라이드 번호와 다를 수 있다)."""
    rels = f"ppt/slides/_rels/{slide_name.rsplit('/', 1)[1]}.rels"
    try:
        root = ET.fromstring(z.read(rels))
    except KeyError:
        return None
    for rel in root:
        if rel.get("Type", "").endswith("/notesSlide"):
            return "ppt/" + rel.get("Target", "").replace("../", "")
    return None


def _pptx(path):
    out = []
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        slides = sorted((n for n in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)), key=_num)
        for i, name in enumerate(slides, 1):
            root = ET.fromstring(z.read(name))
            lines = []
            in_table = set()
            for tbl in root.iter(A + "tbl"):
                for tr in tbl.iter(A + "tr"):
                    cells = []
                    for tc in tr.iter(A + "tc"):
                        paras = list(tc.iter(A + "p"))
                        in_table.update(id(p) for p in paras)
                        cells.append(" ".join(_a_text(p).strip() for p in paras).strip())
                    lines.append("| " + " | ".join(cells) + " |")
            # 표는 위에서 행 단위로 모았으므로, 표 밖 문단을 앞에 둔다.
            text_lines = [_a_text(p) for p in root.iter(A + "p") if id(p) not in in_table and _a_text(p).strip()]
            out.append(f"--- 슬라이드 {i} ---\n" + "\n".join(text_lines + lines))
            notes = _pptx_notes(z, name)
            if notes in names:
                root = ET.fromstring(z.read(notes))
                note = "\n".join("".join(t.text or "" for t in p.iter(A + "t")) for p in root.iter(A + "p")).strip()
                # 노트 슬라이드에는 슬라이드 번호 자리표시자가 숫자로 들어 있다.
                if note and not note.isdigit():
                    out.append("[발표자 노트]\n" + note)
    return "\n\n".join(out)


# ---- Excel ----

def _xlsx(path):
    from openpyxl import load_workbook

    # 수식은 마지막으로 저장된 계산 값으로 읽고, 계산 값이 없으면(Excel에서 열어 저장한 적이 없는 파일) 수식을 보여 준다.
    # 셀 주소를 함께 적어 excel.write_cells로 고칠 수 있게 한다.
    wb = load_workbook(str(path), read_only=True, data_only=True)
    wb_formula = load_workbook(str(path), read_only=True, data_only=False)
    out = []
    try:
        for ws, ws_formula in zip(wb.worksheets, wb_formula.worksheets):
            lines = [f"## 시트: {ws.title}"]
            rows = 0
            for row, row_formula in zip(ws.iter_rows(max_col=MAX_SHEET_COLS), ws_formula.iter_rows(max_col=MAX_SHEET_COLS)):
                cells = []
                for c, f in zip(row, row_formula):
                    value = c.value if c.value is not None else f.value
                    # 읽기 전용 모드의 빈 셀(EmptyCell)에는 coordinate가 없다.
                    if getattr(f, "coordinate", None) and value is not None and str(value).strip() != "":
                        cells.append(f"{f.coordinate}={_cell(value)}")
                if not cells:
                    continue
                lines.append("  ".join(cells))
                rows += 1
                if rows >= MAX_SHEET_ROWS:
                    lines.append(f"... ({MAX_SHEET_ROWS}행까지만 표시)")
                    break
            if rows == 0:
                lines.append("(빈 시트)")
            out.append("\n".join(lines))
    finally:
        wb.close()
        wb_formula.close()
    return "\n\n".join(out)


def _cell(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).replace("\n", " ")


# ---- 한글 ----

def _hwpx(path):
    out = []
    with zipfile.ZipFile(path) as z:
        sections = sorted((n for n in z.namelist() if re.fullmatch(r"Contents/section\d+\.xml", n)), key=_num)
        for name in sections:
            root = ET.fromstring(z.read(name))
            for p in root.iter():
                if p.tag.endswith("}p"):
                    # 표 안 문단은 바깥 문단에도 포함되므로 바로 아래 run의 글자만 모은다.
                    text = "".join((t.text or "") for run in p if run.tag.endswith("}run")
                                   for t in run if t.tag.endswith("}t"))
                    if text.strip():
                        out.append(text)
    return "\n".join(out)


HWP_PARA_TEXT = 67
# 한 글자짜리 제어 문자. 나머지 0~31은 8글자(16바이트)를 차지하는 인라인·확장 제어다.
HWP_CHAR_CONTROLS = {0, 10, 13} | set(range(24, 32))


def _hwp(path):
    import olefile

    if not olefile.isOleFile(str(path)):
        raise DocumentError("HWP 5.0 형식이 아닙니다(HWP 3.x 이하이거나 손상된 파일).")
    with olefile.OleFileIO(str(path)) as ole:
        header = ole.openstream("FileHeader").read()
        flags = struct.unpack("<I", header[36:40])[0]
        if flags & 0x2:
            raise DocumentError("암호가 걸린 HWP 문서입니다.")
        if flags & 0x4:
            raise DocumentError("배포용 HWP 문서는 읽을 수 없습니다.")
        compressed = bool(flags & 0x1)
        sections = sorted((e for e in ole.listdir() if len(e) == 2 and e[0] == "BodyText"),
                          key=lambda e: int(re.sub(r"\D", "", e[1]) or 0))
        out = []
        for entry in sections:
            data = ole.openstream(entry).read()
            if compressed:
                data = zlib.decompress(data, -15)
            out.append(_hwp_section(data))
    return "\n".join(out)


def _hwp_section(data):
    lines, pos = [], 0
    while pos + 4 <= len(data):
        header = struct.unpack("<I", data[pos:pos + 4])[0]
        pos += 4
        tag, size = header & 0x3FF, header >> 20
        if size == 0xFFF:
            size = struct.unpack("<I", data[pos:pos + 4])[0]
            pos += 4
        if tag == HWP_PARA_TEXT:
            lines.append(_hwp_text(data[pos:pos + size]))
        pos += size
    return "\n".join(line for line in lines if line.strip())


def _hwp_text(raw):
    units = struct.unpack(f"<{len(raw) // 2}H", raw[:len(raw) // 2 * 2])
    chars, i = [], 0
    while i < len(units):
        c = units[i]
        if c >= 32:
            chars.append(chr(c))
            i += 1
        elif c in HWP_CHAR_CONTROLS:
            if c == 10:
                chars.append("\n")
            i += 1
        else:
            if c == 9:
                chars.append("\t")
            i += 8
    return "".join(chars)
